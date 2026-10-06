#!/usr/bin/env python3
"""Read-only diagnostics for the BUCAP barcode import."""

from __future__ import annotations

import csv
import json
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook
from sqlalchemy import func

from app import app, db
from app.models import Area, Box, Section, Site, Volume


ROOT = Path(__file__).resolve().parent
INPUT_FILE = ROOT / "bucap.xlsx"
OUTPUT_DIR = ROOT / "bucap_diagnostics"
EXPECTED_HEADER = "BARCODE LABEL NO."
BARCODE_COLUMN = 2
HEADER_ROW = 4
DATA_START_ROW = 5
QUERY_CHUNK_SIZE = 500
VOLUME_SAMPLE_SIZE = 5
AMBIGUOUS_FOCUS = "60979"


def normalise_box_name(value: Any) -> str | None:
    """Use the same Excel-value normalization as update_bucap_boxes.py."""
    if value is None:
        return None
    if isinstance(value, str):
        result = value.strip()
    elif isinstance(value, bool):
        result = str(value)
    elif isinstance(value, int):
        result = str(value)
    elif isinstance(value, float):
        result = str(int(value)) if value.is_integer() else str(value)
    elif isinstance(value, (datetime, date, time)):
        result = value.isoformat()
    else:
        result = str(value).strip()
    return result or None


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def chunks(values: Sequence[Any], size: int = QUERY_CHUNK_SIZE) -> Iterable[Sequence[Any]]:
    for start in range(0, len(values), size):
        yield values[start : start + size]


def unique_headers(raw_headers: Sequence[Any]) -> list[str]:
    headers: list[str] = []
    seen: Counter[str] = Counter()
    for index, value in enumerate(raw_headers, start=1):
        base = str(value).strip() if value is not None and str(value).strip() else f"column_{index}"
        seen[base] += 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return headers


def read_excel(path: Path) -> tuple[list[str], list[dict[str, Any]], dict[str, int]]:
    if not path.is_file():
        raise FileNotFoundError(f"Excel file not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        raw_header = worksheet.cell(row=HEADER_ROW, column=BARCODE_COLUMN).value
        actual_header = raw_header.strip() if isinstance(raw_header, str) else raw_header
        if actual_header != EXPECTED_HEADER:
            raise ValueError(
                f"invalid Excel format: B4 must be {EXPECTED_HEADER!r} after "
                f"trimming; found {raw_header!r}"
            )

        max_column = worksheet.max_column
        raw_headers = [
            worksheet.cell(row=HEADER_ROW, column=column).value
            for column in range(1, max_column + 1)
        ]
        headers = unique_headers(raw_headers)
        occurrences: list[dict[str, Any]] = []
        stats = {"data_rows_examined": 0, "empty_barcodes": 0}

        for excel_row, values in enumerate(
            worksheet.iter_rows(
                min_row=DATA_START_ROW,
                max_col=max_column,
                values_only=True,
            ),
            start=DATA_START_ROW,
        ):
            stats["data_rows_examined"] += 1
            barcode = normalise_box_name(values[BARCODE_COLUMN - 1])
            if barcode is None:
                stats["empty_barcodes"] += 1
                continue
            occurrences.append(
                {
                    "excel_row": excel_row,
                    "box_name": barcode,
                    "original": {
                        header: cell_text(value)
                        for header, value in zip(headers, values)
                    },
                }
            )
        return headers, occurrences, stats
    finally:
        workbook.close()


def load_box_matches(barcodes: list[str]) -> dict[str, list[tuple[int, str, int | None]]]:
    matches: dict[str, list[tuple[int, str, int | None]]] = defaultdict(list)
    for barcode_chunk in chunks(barcodes):
        rows = (
            db.session.query(Box.id, Box.name, Box.section_id)
            .filter(Box.name.in_(barcode_chunk))
            .order_by(Box.name, Box.id)
            .all()
        )
        for box_id, name, section_id in rows:
            matches[str(name)].append((box_id, str(name), section_id))
    return matches


def classify_barcodes(
    barcodes: list[str],
    occurrences_by_barcode: dict[str, list[dict[str, Any]]],
    matches: dict[str, list[tuple[int, str, int | None]]],
) -> list[dict[str, Any]]:
    classification: list[dict[str, Any]] = []
    for barcode in barcodes:
        box_rows = matches.get(barcode, [])
        count = len(box_rows)
        status = "NOT_FOUND" if count == 0 else "FOUND" if count == 1 else "AMBIGUOUS"
        occurrence_rows = occurrences_by_barcode[barcode]
        classification.append(
            {
                "box_name": barcode,
                "status": status,
                "matched_box_count": count,
                "box_ids": ";".join(str(row[0]) for row in box_rows),
                "current_section_ids": ";".join(str(row[2]) for row in box_rows),
                "excel_occurrences": len(occurrence_rows),
                "first_excel_row": occurrence_rows[0]["excel_row"],
            }
        )
    return classification


def search_text_field(
    model: Any,
    column: Any,
    field_label: str,
    barcodes: list[str],
    errors: list[str],
) -> dict[str, list[dict[str, Any]]]:
    local: dict[str, list[dict[str, Any]]] = defaultdict(list)
    try:
        for barcode_chunk in chunks(barcodes):
            rows = (
                db.session.query(model.id, column)
                .filter(column.in_(barcode_chunk))
                .order_by(model.id)
                .all()
            )
            for record_id, value in rows:
                local[str(value)].append(
                    {"field": field_label, "record_id": record_id, "value": cell_text(value)}
                )
    except Exception as exc:
        db.session.rollback()
        errors.append(f"Alternative search {field_label} failed: {exc}")
        return {}
    return local


def search_integer_field(
    model: Any,
    column: Any,
    field_label: str,
    barcodes: list[str],
    errors: list[str],
) -> dict[str, list[dict[str, Any]]]:
    numeric_values: list[int] = []
    canonical_to_originals: dict[str, list[str]] = defaultdict(list)
    for barcode in barcodes:
        try:
            number = int(barcode)
        except ValueError:
            continue
        if str(number) == barcode or barcode.lstrip("0") == str(number):
            numeric_values.append(number)
            canonical_to_originals[str(number)].append(barcode)

    local: dict[str, list[dict[str, Any]]] = defaultdict(list)
    try:
        unique_numbers = list(dict.fromkeys(numeric_values))
        for number_chunk in chunks(unique_numbers):
            rows = (
                db.session.query(model.id, column)
                .filter(column.in_(number_chunk))
                .order_by(model.id)
                .all()
            )
            for record_id, value in rows:
                for original in canonical_to_originals[str(value)]:
                    local[original].append(
                        {
                            "field": field_label,
                            "record_id": record_id,
                            "value": cell_text(value),
                        }
                    )
    except Exception as exc:
        db.session.rollback()
        errors.append(f"Alternative search {field_label} failed: {exc}")
        return {}
    return local


def alternative_searches(
    not_found: list[str], errors: list[str]
) -> dict[str, list[dict[str, Any]]]:
    combined: dict[str, list[dict[str, Any]]] = defaultdict(list)
    searches = [
        search_text_field(Volume, Volume.name, "Volume.name", not_found, errors),
        search_text_field(
            Volume, Volume.order_number, "Volume.order_number", not_found, errors
        ),
        search_integer_field(
            Volume, Volume.account_number, "Volume.account_number", not_found, errors
        ),
    ]
    for result in searches:
        for barcode, entries in result.items():
            combined[barcode].extend(entries)
    return combined


def load_ambiguous_details(
    ambiguous_box_ids: list[int], errors: list[str]
) -> list[dict[str, Any]]:
    if not ambiguous_box_ids:
        return []
    try:
        hierarchy_rows = (
            db.session.query(
                Box.id,
                Box.name,
                Box.section_id,
                Section.name,
                Area.id,
                Area.name,
                Site.id,
                Site.name,
            )
            .outerjoin(Section, Box.section_id == Section.id)
            .outerjoin(Area, Section.area_id == Area.id)
            .outerjoin(Site, Area.site_id == Site.id)
            .filter(Box.id.in_(ambiguous_box_ids))
            .order_by(Box.name, Box.id)
            .all()
        )
        volume_counts = dict(
            db.session.query(Volume.box_id, func.count(Volume.id))
            .filter(Volume.box_id.in_(ambiguous_box_ids))
            .group_by(Volume.box_id)
            .all()
        )
        samples: dict[int, list[str]] = defaultdict(list)
        volume_rows = (
            db.session.query(Volume.box_id, Volume.name)
            .filter(Volume.box_id.in_(ambiguous_box_ids))
            .order_by(Volume.box_id, Volume.id)
            .all()
        )
        for box_id, volume_name in volume_rows:
            if len(samples[box_id]) < VOLUME_SAMPLE_SIZE:
                samples[box_id].append(cell_text(volume_name))

        details = []
        for row in hierarchy_rows:
            details.append(
                {
                    "box_id": row[0],
                    "box_name": row[1],
                    "section_id": row[2],
                    "section_name": row[3],
                    "area_id": row[4],
                    "area_name": row[5],
                    "site_id": row[6],
                    "site_name": row[7],
                    "volume_count": volume_counts.get(row[0], 0),
                    "volume_name_sample": " | ".join(samples.get(row[0], [])),
                }
            )
        return details
    except Exception as exc:
        db.session.rollback()
        errors.append(f"Ambiguous hierarchy diagnosis failed: {exc}")
        return []


def consecutive_groups(numbers: list[int], maximum_gap: int) -> list[list[int]]:
    if not numbers:
        return []
    groups: list[list[int]] = [[numbers[0]]]
    for number in numbers[1:]:
        if number - groups[-1][-1] <= maximum_gap:
            groups[-1].append(number)
        else:
            groups.append([number])
    return [group for group in groups if len(group) > 1]


def row_groups(rows: list[int], maximum_gap: int) -> list[list[int]]:
    return consecutive_groups(sorted(set(rows)), maximum_gap)


def format_groups(groups: list[list[int]], limit: int = 100) -> list[str]:
    lines = []
    for group in sorted(groups, key=lambda values: (-len(values), values[0]))[:limit]:
        lines.append(f"{group[0]}-{group[-1]} ({len(group)} values)")
    return lines


def category_comparison(
    headers: list[str],
    occurrences: list[dict[str, Any]],
    status_by_barcode: dict[str, str],
) -> list[str]:
    barcode_header = headers[BARCODE_COLUMN - 1]
    lines: list[str] = []
    for header in headers:
        if header == barcode_header:
            continue
        found: Counter[str] = Counter()
        not_found: Counter[str] = Counter()
        for occurrence in occurrences:
            value = occurrence["original"].get(header, "") or "<EMPTY>"
            status = status_by_barcode[occurrence["box_name"]]
            if status == "NOT_FOUND":
                not_found[value] += 1
            elif status == "FOUND":
                found[value] += 1
        distinct = len(set(found) | set(not_found))
        lines.append(
            f"Column {header!r}: {distinct} distinct values; "
            f"FOUND occurrences={sum(found.values())}; "
            f"NOT_FOUND occurrences={sum(not_found.values())}"
        )
        for value, nf_count in not_found.most_common(15):
            found_count = found.get(value, 0)
            total = nf_count + found_count
            lines.append(
                f"  {value!r}: NOT_FOUND={nf_count}, FOUND={found_count}, "
                f"not_found_share={nf_count / total:.1%}"
            )
    return lines


def build_summary(
    headers: list[str],
    occurrences: list[dict[str, Any]],
    excel_stats: dict[str, int],
    classification: list[dict[str, Any]],
    alternative_matches: dict[str, list[dict[str, Any]]],
    ambiguous_details: list[dict[str, Any]],
    errors: list[str],
) -> str:
    status_counts = Counter(row["status"] for row in classification)
    status_by_barcode = {row["box_name"]: row["status"] for row in classification}
    not_found_names = [
        row["box_name"] for row in classification if row["status"] == "NOT_FOUND"
    ]
    numeric_pairs: list[tuple[int, str]] = []
    for barcode in not_found_names:
        try:
            numeric_pairs.append((int(barcode), barcode))
        except ValueError:
            pass
    numeric_values = sorted(set(pair[0] for pair in numeric_pairs))
    prefix2 = Counter(barcode[:2] for barcode in not_found_names)
    prefix3 = Counter(barcode[:3] for barcode in not_found_names)
    not_found_rows = [
        occurrence["excel_row"]
        for occurrence in occurrences
        if status_by_barcode[occurrence["box_name"]] == "NOT_FOUND"
    ]

    ordered_statuses = [status_by_barcode[item["box_name"]] for item in occurrences]
    transitions = Counter(
        f"{left}->{right}"
        for left, right in zip(ordered_statuses, ordered_statuses[1:])
        if left != right
    )
    alternative_field_counts = Counter(
        match["field"]
        for barcode in not_found_names
        for match in alternative_matches.get(barcode, [])
    )

    lines = [
        "BUCAP BOX DIAGNOSTIC SUMMARY",
        "",
        "GENERAL COUNTS",
        f"Data rows examined: {excel_stats['data_rows_examined']}",
        f"Empty barcode cells: {excel_stats['empty_barcodes']}",
        f"Unique barcodes: {len(classification)}",
        f"FOUND: {status_counts['FOUND']}",
        f"NOT_FOUND: {status_counts['NOT_FOUND']}",
        f"AMBIGUOUS: {status_counts['AMBIGUOUS']}",
        f"NOT_FOUND present in alternative fields: "
        f"{sum(bool(alternative_matches.get(name)) for name in not_found_names)}",
        "",
        "ALTERNATIVE EXACT-MATCH SEARCHES",
        "Fields: Volume.name, Volume.order_number, Volume.account_number",
    ]
    if alternative_field_counts:
        lines.extend(
            f"{field}: {count} matching database records"
            for field, count in alternative_field_counts.most_common()
        )
    else:
        lines.append("No alternative exact matches found.")

    lines.extend(["", "NOT_FOUND NUMERIC PATTERNS"])
    lines.append(f"Numerically convertible: {len(numeric_pairs)}")
    lines.append(f"Numeric minimum: {numeric_values[0] if numeric_values else 'N/A'}")
    lines.append(f"Numeric maximum: {numeric_values[-1] if numeric_values else 'N/A'}")
    lines.append("First 2 digits distribution:")
    lines.extend(f"  {prefix}: {count}" for prefix, count in prefix2.most_common())
    lines.append("First 3 digits distribution:")
    lines.extend(f"  {prefix}: {count}" for prefix, count in prefix3.most_common())
    lines.append("Consecutive numeric ranges (gap <= 1):")
    lines.extend(f"  {line}" for line in format_groups(consecutive_groups(numeric_values, 1)))
    lines.append("Nearby numeric clusters (gap <= 5):")
    lines.extend(f"  {line}" for line in format_groups(consecutive_groups(numeric_values, 5)))

    lines.extend(["", "EXCEL ROW POSITION PATTERNS"])
    lines.append("Consecutive NOT_FOUND row blocks (gap <= 1):")
    lines.extend(f"  {line}" for line in format_groups(row_groups(not_found_rows, 1)))
    lines.append("Nearby NOT_FOUND row groups (gap <= 3):")
    lines.extend(f"  {line}" for line in format_groups(row_groups(not_found_rows, 3)))
    lines.append("Status transitions between adjacent data occurrences:")
    lines.extend(f"  {transition}: {count}" for transition, count in transitions.most_common())

    lines.extend(["", "FOUND VS NOT_FOUND BY EXCEL COLUMN"])
    lines.extend(category_comparison(headers, occurrences, status_by_barcode))

    lines.extend(["", f"AMBIGUOUS FOCUS: Box.name={AMBIGUOUS_FOCUS!r}"])
    focus_rows = [row for row in ambiguous_details if row["box_name"] == AMBIGUOUS_FOCUS]
    if not focus_rows:
        lines.append("No detailed records available for the requested focus barcode.")
    for row in focus_rows:
        lines.append(
            f"Box.id={row['box_id']}; section_id={row['section_id']}; "
            f"Section.name={row['section_name']!r}; area_id={row['area_id']}; "
            f"Area.name={row['area_name']!r}; site_id={row['site_id']}; "
            f"Site.name={row['site_name']!r}; volumes={row['volume_count']}; "
            f"Volume.name sample={row['volume_name_sample']!r}"
        )

    lines.extend(["", "DIAGNOSTIC ERRORS"])
    lines.extend(errors or ["None"])
    return "\n".join(lines) + "\n"


def write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_reports(
    headers: list[str],
    occurrences: list[dict[str, Any]],
    excel_stats: dict[str, int],
    classification: list[dict[str, Any]],
    alternative_matches: dict[str, list[dict[str, Any]]],
    ambiguous_details: list[dict[str, Any]],
    errors: list[str],
) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    classification_fields = [
        "box_name",
        "status",
        "matched_box_count",
        "box_ids",
        "current_section_ids",
        "excel_occurrences",
        "first_excel_row",
    ]
    write_csv(
        OUTPUT_DIR / "bucap_box_classification.csv",
        classification_fields,
        classification,
    )

    status_by_barcode = {row["box_name"]: row["status"] for row in classification}
    not_found_rows = []
    for occurrence in occurrences:
        barcode = occurrence["box_name"]
        if status_by_barcode[barcode] != "NOT_FOUND":
            continue
        matches = alternative_matches.get(barcode, [])
        row = {
            "excel_row": occurrence["excel_row"],
            "box_name": barcode,
            "alternative_match_count": len(matches),
            "alternative_db_matches": json.dumps(matches, ensure_ascii=False),
        }
        row.update(
            {f"excel_{header}": value for header, value in occurrence["original"].items()}
        )
        not_found_rows.append(row)
    not_found_fields = [
        "excel_row",
        "box_name",
        *[f"excel_{header}" for header in headers],
        "alternative_match_count",
        "alternative_db_matches",
    ]
    write_csv(
        OUTPUT_DIR / "bucap_not_found.csv", not_found_fields, not_found_rows
    )

    ambiguous_fields = [
        "box_id",
        "box_name",
        "section_id",
        "section_name",
        "area_id",
        "area_name",
        "site_id",
        "site_name",
        "volume_count",
        "volume_name_sample",
    ]
    write_csv(
        OUTPUT_DIR / "bucap_ambiguous.csv", ambiguous_fields, ambiguous_details
    )

    summary = build_summary(
        headers,
        occurrences,
        excel_stats,
        classification,
        alternative_matches,
        ambiguous_details,
        errors,
    )
    (OUTPUT_DIR / "bucap_diagnostic_summary.txt").write_text(
        summary, encoding="utf-8"
    )


def main() -> int:
    errors: list[str] = []
    print(f"Excel file: {INPUT_FILE}")
    try:
        headers, occurrences, excel_stats = read_excel(INPUT_FILE)
        occurrences_by_barcode: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for occurrence in occurrences:
            occurrences_by_barcode[occurrence["box_name"]].append(occurrence)
        barcodes = list(occurrences_by_barcode)

        with app.app_context():
            try:
                matches = load_box_matches(barcodes)
                classification = classify_barcodes(
                    barcodes, occurrences_by_barcode, matches
                )
                not_found = [
                    row["box_name"]
                    for row in classification
                    if row["status"] == "NOT_FOUND"
                ]
                alternative_matches = alternative_searches(not_found, errors)
                ambiguous_ids = [
                    box_id
                    for row in classification
                    if row["status"] == "AMBIGUOUS"
                    for box_id in (
                        int(value) for value in row["box_ids"].split(";") if value
                    )
                ]
                ambiguous_details = load_ambiguous_details(ambiguous_ids, errors)
                write_reports(
                    headers,
                    occurrences,
                    excel_stats,
                    classification,
                    alternative_matches,
                    ambiguous_details,
                    errors,
                )

                counts = Counter(row["status"] for row in classification)
                alternative_count = sum(
                    bool(alternative_matches.get(barcode)) for barcode in not_found
                )
                print(f"Unique barcodes: {len(classification)}")
                print(f"FOUND: {counts['FOUND']}")
                print(f"NOT_FOUND: {counts['NOT_FOUND']}")
                print(f"AMBIGUOUS: {counts['AMBIGUOUS']}")
                print(f"NOT_FOUND matched in other DB fields: {alternative_count}")
                print(f"Reports directory: {OUTPUT_DIR}")
                print(f"Diagnostic errors: {len(errors)}")
                for error in errors:
                    print(f"ERROR: {error}", file=sys.stderr)
                return 0 if not errors else 1
            finally:
                db.session.rollback()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
