#!/usr/bin/env python3
"""Move the boxes listed in bucap.xlsx to the configured ADM section.

The default mode is a dry run. Use --commit explicitly to persist changes.
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app import app, db
from app.models import Area, Box, Section, Site, Volume


SITE_ID = 52
AREA_ID = 61
SECTION_ID = 26684
DEFAULT_INPUT = Path(__file__).resolve().parent / "bucap.xlsx"
EXPECTED_HEADER = "BARCODE LABEL NO."


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Move the boxes in the second Excel column to section 26684."
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--dry-run",
        action="store_true",
        help="show the planned changes without committing (default)",
    )
    mode.add_argument(
        "--commit",
        action="store_true",
        help="apply all changes in one database transaction",
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT,
        help=f"Excel input file (default: {DEFAULT_INPUT.name})",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="print the classification details for every unique Box.name",
    )
    return parser.parse_args()


def normalise_box_name(value: Any) -> str | None:
    """Convert an Excel value to the string used for Box.name."""
    if value is None:
        return None
    if isinstance(value, str):
        result = value.strip()
    elif isinstance(value, bool):
        result = str(value)
    elif isinstance(value, int):
        result = str(value)
    elif isinstance(value, float):
        result = str(int(value)) if value.is_integer() else str(value)
    elif isinstance(value, (datetime, date, time)):
        result = value.isoformat()
    else:
        result = str(value).strip()
    return result or None


def read_box_names(path: Path) -> tuple[list[str], dict[str, int]]:
    if not path.is_file():
        raise FileNotFoundError(f"Excel file not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        raw_header = worksheet["B4"].value
        actual_header = raw_header.strip() if isinstance(raw_header, str) else raw_header
        if actual_header != EXPECTED_HEADER:
            raise ValueError(
                "invalid Excel format: cell B4 must contain exactly "
                f"{EXPECTED_HEADER!r} after trimming; found {raw_header!r}"
            )

        names: list[str] = []
        stats = {"data_rows_examined": 0, "empty": 0}

        for row in worksheet.iter_rows(
            min_row=5, min_col=2, max_col=2, values_only=True
        ):
            stats["data_rows_examined"] += 1
            name = normalise_box_name(row[0])
            if name is None:
                stats["empty"] += 1
                continue
            names.append(name)

        return names, stats
    finally:
        workbook.close()


def analyse_target_hierarchy() -> tuple[Site | None, Area | None, Section | None]:
    site = db.session.query(Site).filter(Site.id == SITE_ID).one_or_none()
    area = db.session.query(Area).filter(Area.id == AREA_ID).one_or_none()
    section = db.session.query(Section).filter(Section.id == SECTION_ID).one_or_none()
    return site, area, section


def print_summary(summary: dict[str, Any]) -> None:
    print("\nEXCEL / BOX SUMMARY")
    print(f"  Data rows examined:    {summary['data_rows_examined']}")
    print(f"  Empty cells:           {summary['empty']}")
    print(f"  Valid unique boxes:    {summary['valid']}")
    print(f"  Duplicate rows:        {summary['duplicates']}")
    print(f"  Already correct:       {summary['already_correct']}")
    print(f"  Boxes to modify:       {summary['to_modify']}")
    print(f"  Boxes not found:       {summary['not_found']}")
    print(f"  Ambiguous boxes:       {summary['ambiguous']}")
    print(f"  Errors:                {summary['errors']}")


def main() -> int:
    args = parse_args()
    commit_mode = bool(args.commit)
    mode_name = "COMMIT" if commit_mode else "DRY RUN"
    summary = {
        "data_rows_examined": 0,
        "empty": 0,
        "valid": 0,
        "already_correct": 0,
        "to_modify": 0,
        "not_found": 0,
        "ambiguous": 0,
        "duplicates": 0,
        "errors": 0,
        "site_status": "NOT CHECKED",
        "area_status": "NOT CHECKED",
        "section_status": "NOT CHECKED",
    }

    print(f"Mode: {mode_name}")
    print(f"Input: {args.input.resolve()}")

    try:
        names, excel_stats = read_box_names(args.input.resolve())
        summary.update(excel_stats)

        counts = Counter(names)
        unique_names = list(dict.fromkeys(names))
        summary["duplicates"] = sum(count - 1 for count in counts.values())
        summary["valid"] = len(unique_names)

        if args.verbose and summary["duplicates"]:
            print("\nDUPLICATES IN EXCEL")
            for name, count in counts.items():
                if count > 1:
                    print(f"  {name!r}: {count} occurrences")

        with app.app_context():
            site, area, section = analyse_target_hierarchy()
            summary["site_status"] = "EXISTS" if site is not None else "WOULD CREATE"
            summary["area_status"] = "EXISTS" if area is not None else "WOULD CREATE"
            summary["section_status"] = (
                "EXISTS" if section is not None else "WOULD CREATE"
            )

            if area is not None and area.site_id != SITE_ID:
                raise RuntimeError(
                    f"hierarchy collision: Area.id={AREA_ID} has "
                    f"site_id={area.site_id}; expected site_id={SITE_ID}. "
                    "Existing data will not be changed."
                )
            if section is not None and section.area_id != AREA_ID:
                raise RuntimeError(
                    f"hierarchy collision: Section.id={SECTION_ID} has "
                    f"area_id={section.area_id}; expected area_id={AREA_ID}. "
                    "Existing data will not be changed."
                )

            print("\nTARGET HIERARCHY")
            if site is None:
                print(f"  WOULD CREATE | Site id={SITE_ID}, name='BUCAP'")
            else:
                print(f"  EXISTS | Site id={site.id}, name={site.name!r}")
            if area is None:
                print(
                    f"  WOULD CREATE | Area id={AREA_ID}, name='BUCAP', "
                    f"site_id={SITE_ID}"
                )
            else:
                print(
                    f"  EXISTS | Area id={area.id}, name={area.name!r}, "
                    f"site_id={area.site_id}"
                )
            if section is None:
                print(
                    f"  WOULD CREATE | Section id={SECTION_ID}, name='BUCAP', "
                    f"area_id={AREA_ID}"
                )
            else:
                print(
                    f"  EXISTS | Section id={section.id}, name={section.name!r}, "
                    f"area_id={section.area_id}"
                )

            boxes_to_modify: list[Box] = []
            affected_box_ids: list[int] = []
            not_found_names: list[str] = []
            ambiguous_boxes: list[tuple[str, list[Box]]] = []
            current_section_distribution: Counter[int | None] = Counter()
            if args.verbose:
                print("\nBOX CHECK")
            for name in unique_names:
                matching_boxes = (
                    db.session.query(Box)
                    .filter(Box.name == name)
                    .order_by(Box.id)
                    .all()
                )
                if not matching_boxes:
                    summary["not_found"] += 1
                    not_found_names.append(name)
                    continue
                if len(matching_boxes) > 1:
                    summary["ambiguous"] += 1
                    ambiguous_boxes.append((name, matching_boxes))
                    continue

                box = matching_boxes[0]
                affected_box_ids.append(box.id)
                if box.section_id == SECTION_ID:
                    summary["already_correct"] += 1
                    status = "ALREADY CORRECT"
                else:
                    summary["to_modify"] += 1
                    boxes_to_modify.append(box)
                    current_section_distribution[box.section_id] += 1
                    status = "TO MODIFY"
                if args.verbose:
                    print(
                        f"  {status} | name={box.name!r} | id={box.id} | "
                        f"current_section_id={box.section_id} | "
                        f"target_section_id={SECTION_ID}"
                    )

            if not_found_names:
                print(f"\nNOT FOUND BOXES: {len(not_found_names)}")
                if args.verbose:
                    for name in not_found_names:
                        print(f"  {name}")
                else:
                    preview = not_found_names[:20]
                    print(f"  First {len(preview)}: {', '.join(preview)}")
                    additional = len(not_found_names) - len(preview)
                    if additional:
                        print(f"  ... {additional} additional NOT FOUND boxes")

            if ambiguous_boxes:
                print("\nAMBIGUOUS BOXES")
                for name, matching_boxes in ambiguous_boxes:
                    details = "; ".join(
                        f"id={box.id} section_id={box.section_id}"
                        for box in matching_boxes
                    )
                    print(f"  name={name!r} -> {details}")

            if args.verbose:
                print("\nCURRENT SECTION DISTRIBUTION")
                if current_section_distribution:
                    for current_section_id, count in sorted(
                        current_section_distribution.items(),
                        key=lambda item: (-item[1], item[0] is None, item[0]),
                    ):
                        print(f"  section_id={current_section_id}: {count}")
                else:
                    print("  No boxes to modify.")

            classified = (
                summary["already_correct"]
                + summary["to_modify"]
                + summary["not_found"]
                + summary["ambiguous"]
            )
            if classified != summary["valid"]:
                raise RuntimeError(
                    "internal classification check failed: already correct + "
                    "to modify + not found + ambiguous does not equal valid "
                    "unique boxes"
                )

            # This snapshot makes the invariant explicit: moving a Box must not
            # alter any Volume.box_id association.
            volume_snapshot = (
                db.session.query(Volume.id, Volume.box_id)
                .filter(Volume.box_id.in_(affected_box_ids))
                .order_by(Volume.id)
                .all()
                if affected_box_ids
                else []
            )

            if commit_mode:
                if summary["not_found"] or summary["ambiguous"]:
                    raise RuntimeError(
                        f"commit refused: {summary['not_found']} unique Box.name "
                        "value(s) were not found and "
                        f"{summary['ambiguous']} were ambiguous"
                    )

                # Mutations start only after the complete hierarchy and Box analysis.
                create_site = site is None
                create_area = area is None
                create_section = section is None
                if site is None:
                    # country, city and address are mandatory in the real Site model.
                    site = Site(
                        id=SITE_ID,
                        name="BUCAP",
                        country="N/A",
                        city="N/A",
                        address="N/A",
                    )
                    db.session.add(site)
                    db.session.flush()

                if area is None:
                    area = Area(id=AREA_ID, name="BUCAP", site_id=SITE_ID)
                    db.session.add(area)
                    db.session.flush()

                if section is None:
                    section = Section(
                        id=SECTION_ID, name="BUCAP", area_id=AREA_ID
                    )
                    db.session.add(section)
                    db.session.flush()

                for box in boxes_to_modify:
                    box.section_id = SECTION_ID

                db.session.flush()
                volume_after = (
                    db.session.query(Volume.id, Volume.box_id)
                    .filter(Volume.box_id.in_(affected_box_ids))
                    .order_by(Volume.id)
                    .all()
                    if affected_box_ids
                    else []
                )
                if volume_after != volume_snapshot:
                    raise RuntimeError("Volume.box_id invariant check failed")

                db.session.commit()
                if create_site:
                    summary["site_status"] = "CREATED"
                if create_area:
                    summary["area_status"] = "CREATED"
                if create_section:
                    summary["section_status"] = "CREATED"
                print(f"\nCOMMITTED: {len(boxes_to_modify)} box(es) updated.")
            else:
                db.session.rollback()
                print("\nDRY RUN: no changes were committed.")

    except Exception as exc:
        summary["errors"] += 1
        try:
            with app.app_context():
                db.session.rollback()
        except Exception:
            pass
        print(f"\nERROR: {exc}", file=sys.stderr)
        print("No database changes were committed.", file=sys.stderr)
        print_summary(summary)
        return 1
    finally:
        try:
            db.session.remove()
        except Exception:
            pass

    print_summary(summary)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
